from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
import os

# Create a presentation object
prs = Presentation()

# Define slide layouts
TITLE_SLIDE_LAYOUT = 0
BULLET_SLIDE_LAYOUT = 1
PICTURE_SLIDE_LAYOUT = 8

# Slide 1: Title
slide1 = prs.slides.add_slide(prs.slide_layouts[TITLE_SLIDE_LAYOUT])
title = slide1.shapes.title
subtitle = slide1.placeholders[1]
title.text = "IIC AWARENESS PROGRAM"
title.text_frame.paragraphs[0].font.bold = True
title.text_frame.paragraphs[0].font.size = Pt(44)
subtitle.text = "ACE Engineering College\nEmpowering Student Innovators & Entrepreneurs\n\n[Insert ACE Logo Here]"

# Slide 2: What is IIC
slide2 = prs.slides.add_slide(prs.slide_layouts[BULLET_SLIDE_LAYOUT])
shapes2 = slide2.shapes
title_shape2 = shapes2.title
body_shape2 = shapes2.placeholders[1]
title_shape2.text = "What is IIC?"
tf2 = body_shape2.text_frame
tf2.text = "Institution's Innovation Council"
p = tf2.add_paragraph()
p.text = "An initiative by the Ministry of Education's Innovation Cell (MIC) and AICTE."
p = tf2.add_paragraph()
p.text = "A bridge between engineering students and the startup ecosystem."
p = tf2.add_paragraph()
p.text = "Platform to build real-world products, get mentorship, and turn projects into startups!"

# Slide 3: President Murali Sir
slide3 = prs.slides.add_slide(prs.slide_layouts[BULLET_SLIDE_LAYOUT])
shapes3 = slide3.shapes
title_shape3 = shapes3.title
body_shape3 = shapes3.placeholders[1]
title_shape3.text = "President: Dr. M. Murali Sir"
tf3 = body_shape3.text_frame
tf3.text = "Dr. Malijeddi Murali\nVice-Principal and Dean of Skill Development & Industry Integration"
p = tf3.add_paragraph()
p.text = "[Insert Photo Here]"
p = tf3.add_paragraph()
p.text = "Under his leadership, the IIC at ACEEC drives continuous skill development, connecting our academic learning with industry requirements."

# Slide 4: Core Framework & Calendar
slide4 = prs.slides.add_slide(prs.slide_layouts[BULLET_SLIDE_LAYOUT])
shapes4 = slide4.shapes
title_shape4 = shapes4.title
body_shape4 = shapes4.placeholders[1]
title_shape4.text = "Our Core Framework & Calendar"
tf4 = body_shape4.text_frame
tf4.text = "We follow a structured, year-round approach to innovation."
p = tf4.add_paragraph()
p.text = "[Insert IIC 8.0 Calendar Screenshot Here]"

# Slide 5: Student Roles
slide5 = prs.slides.add_slide(prs.slide_layouts[BULLET_SLIDE_LAYOUT])
shapes5 = slide5.shapes
title_shape5 = shapes5.title
body_shape5 = shapes5.placeholders[1]
title_shape5.text = "Student Roles & Opportunities"
tf5 = body_shape5.text_frame
tf5.text = "[Space reserved for student roles - To be updated]"

# Slide 6: PALS Connection
slide6 = prs.slides.add_slide(prs.slide_layouts[PICTURE_SLIDE_LAYOUT])
title_shape6 = slide6.shapes.title
body_shape6 = slide6.placeholders[2]
title_shape6.text = "Industry Connection: PALS"
tf6 = body_shape6.text_frame
tf6.text = "Our college is proudly connected with PALS (Pan IIT Alumni Leadership Series).\n\nBenefits:"
p = tf6.add_paragraph()
p.text = "Mentorship from IIT Alumni and Industry Experts"
p = tf6.add_paragraph()
p.text = "Real-world Innovation Challenges (like innoWAH!)"
p = tf6.add_paragraph()
p.text = "Industry Visits and Internships"

img_path6 = r"C:\Users\prsnl\.gemini\antigravity\brain\d1b69639-1940-4d3a-9d7d-b4963cb4837c\pals_connection_1784021731901.png"
if os.path.exists(img_path6):
    placeholder = slide6.placeholders[1]
    slide6.shapes.add_picture(img_path6, placeholder.left, placeholder.top, placeholder.width, placeholder.height)

# Slide 7: Benefits
slide7 = prs.slides.add_slide(prs.slide_layouts[PICTURE_SLIDE_LAYOUT])
title_shape7 = slide7.shapes.title
body_shape7 = slide7.placeholders[2]
title_shape7.text = "Why Join IIC? Benefits to Students"
tf7 = body_shape7.text_frame
tf7.text = "If you have an idea, we have the resources!"
p = tf7.add_paragraph()
p.text = "From Job Seeker to Job Creator."
p = tf7.add_paragraph()
p.text = "Funding & Incubation via YUKTI Portal."
p = tf7.add_paragraph()
p.text = "Skill Development (Design Thinking, Pitching)."

img_path7 = r"C:\Users\prsnl\.gemini\antigravity\brain\d1b69639-1940-4d3a-9d7d-b4963cb4837c\iic_benefits_1784021742071.png"
if os.path.exists(img_path7):
    placeholder = slide7.placeholders[1]
    slide6.shapes.add_picture(img_path7, placeholder.left, placeholder.top, placeholder.width, placeholder.height)

# Slide 8: Roadmap
slide8 = prs.slides.add_slide(prs.slide_layouts[PICTURE_SLIDE_LAYOUT])
title_shape8 = slide8.shapes.title
body_shape8 = slide8.placeholders[2]
title_shape8.text = "The Roadmap to Success"
tf8 = body_shape8.text_frame
tf8.text = "1. Ideation\n2. Prototyping\n3. Mentorship\n4. Funding\n5. Launch!"

img_path8 = r"C:\Users\prsnl\.gemini\antigravity\brain\d1b69639-1940-4d3a-9d7d-b4963cb4837c\startup_roadmap_1784021754024.png"
if os.path.exists(img_path8):
    placeholder = slide8.placeholders[1]
    slide6.shapes.add_picture(img_path8, placeholder.left, placeholder.top, placeholder.width, placeholder.height)

# Slide 9: IPR
slide9 = prs.slides.add_slide(prs.slide_layouts[BULLET_SLIDE_LAYOUT])
title_shape9 = slide9.shapes.title
body_shape9 = slide9.placeholders[1]
title_shape9.text = "Intellectual Property Rights (IPR)"
tf9 = body_shape9.text_frame
tf9.text = "Protecting your code, your designs, and your brand is crucial."
p = tf9.add_paragraph()
p.text = "Patents: Protect your unique inventions."
p = tf9.add_paragraph()
p.text = "Copyrights: Protect your code and literature."
p = tf9.add_paragraph()
p.text = "Trademarks: Protect your startup's brand name."

# Slide 10: Activities
slide10 = prs.slides.add_slide(prs.slide_layouts[BULLET_SLIDE_LAYOUT])
title_shape10 = slide10.shapes.title
body_shape10 = slide10.placeholders[1]
title_shape10.text = "Major Activities & Events"
tf10 = body_shape10.text_frame
tf10.text = "Startup Bootcamps & Hackathons"
p = tf10.add_paragraph()
p.text = "Prototyping Workshops & Design Thinking"
p = tf10.add_paragraph()
p.text = "Alumni Talks & Expert Mentoring Sessions"
p = tf10.add_paragraph()
p.text = "Idea-Café & R&D Cell Events"

# Slide 11: Real-Life Stories
slide11 = prs.slides.add_slide(prs.slide_layouts[BULLET_SLIDE_LAYOUT])
title_shape11 = slide11.shapes.title
body_shape11 = slide11.placeholders[1]
title_shape11.text = "Real-Life Student Success Stories"
tf11 = body_shape11.text_frame
tf11.text = "The Campus-Grown Startup: Students who started with a final year project, used IIC labs to build a prototype, and launched a SaaS platform."
p = tf11.add_paragraph()
p.text = "The Hackathon Winners: Teams who competed in the Smart India Hackathon (SIH), received seed funding, and are now full-time founders."
p = tf11.add_paragraph()
p.text = "[Insert Photos of Young Founders Here]"

# Slide 12: Govt Support
slide12 = prs.slides.add_slide(prs.slide_layouts[BULLET_SLIDE_LAYOUT])
title_shape12 = slide12.shapes.title
body_shape12 = slide12.placeholders[1]
title_shape12.text = "Backed by the Government"
tf12 = body_shape12.text_frame
tf12.text = "The startup ecosystem in India is heavily supported by:"
p = tf12.add_paragraph()
p.text = "AICTE (All India Council for Technical Education)"
p = tf12.add_paragraph()
p.text = "Ministry of Education's Innovation Cell (MIC)"
p = tf12.add_paragraph()
p.text = "Startup India & Make In India initiatives"

# Slide 13: Vision
slide13 = prs.slides.add_slide(prs.slide_layouts[BULLET_SLIDE_LAYOUT])
title_shape13 = slide13.shapes.title
body_shape13 = slide13.placeholders[1]
title_shape13.text = "Our Vision: Aiming for 5-Stars!"
tf13 = body_shape13.text_frame
tf13.text = "Our ultimate goal is to achieve the prestigious 5-Star Rating for ACE Engineering College's IIC!"
p = tf13.add_paragraph()
p.text = "By participating, you don't just build your own career, you elevate the reputation of our entire college."

# Slide 14: Conclusion
slide14 = prs.slides.add_slide(prs.slide_layouts[BULLET_SLIDE_LAYOUT])
title_shape14 = slide14.shapes.title
body_shape14 = slide14.placeholders[1]
title_shape14.text = "Conclusion"
tf14 = body_shape14.text_frame
tf14.text = "Don't wait until graduation to start building your future. The time to innovate is now."
p = tf14.add_paragraph()
p.text = "Let's build the next big thing, together!"
p = tf14.add_paragraph()
p.text = "[Insert High Energy Images Here]"

# Slide 15: Questions
slide15 = prs.slides.add_slide(prs.slide_layouts[TITLE_SLIDE_LAYOUT])
title15 = slide15.shapes.title
subtitle15 = slide15.placeholders[1]
title15.text = "Any Questions?"
subtitle15.text = "Join the ACE IIC Students Council today!\nContact Dr. Kondal Rao Sir or the Student Ambassadors\nSubmit ideas on: yukti.mic.gov.in"

# Save the presentation
output_path = r"C:\Users\prsnl\Downloads\IIC_Awareness_Presentation.pptx"
prs.save(output_path)
print(f"Presentation saved successfully at {output_path}")
